"""Erstellt Stundenplan.xlsx als Vorlage zum Ausfullen."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DAYS = ["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag"]
HEADERS = ["Modul", "ECTS", "Pruefung (YYYY-MM-DD)", "Schwierigkeit (1-5)"] + DAYS

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Stundenplan"

header_fill = PatternFill("solid", fgColor="2E4057")
header_font = Font(color="FFFFFF", bold=True)
day_fill    = PatternFill("solid", fgColor="048A81")
thin        = Side(style="thin", color="AAAAAA")
border      = Border(left=thin, right=thin, top=thin, bottom=thin)

for col, header in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font      = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = border
    if header in DAYS:
        cell.fill = day_fill
    else:
        cell.fill = header_fill

# Beispielzeilen
examples = [
    ["Statistik",          5, "2026-07-23", 4, "VL", "",   "GUe", "",    ""   ],
    ["Operations Research",5, "2026-07-23", 3, "VL", "",   "",    "GUe", ""   ],
    ["Fertigungstechnik",  4, "2026-08-03", 3, "",   "VL", "",    "",    "GUe"],
]

example_fill = PatternFill("solid", fgColor="F0F4F8")
for row_idx, row_data in enumerate(examples, 2):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.fill      = example_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border    = border

# Spaltenbreiten
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 6
ws.column_dimensions["C"].width = 20
ws.column_dimensions["D"].width = 18
for col in ["E", "F", "G", "H", "I"]:
    ws.column_dimensions[col].width = 12

ws.row_dimensions[1].height = 30

wb.save("Stundenplan.xlsx")
print("Stundenplan.xlsx erstellt.")
print("Trage deine Kurse ein und starte dann: python main.py")
