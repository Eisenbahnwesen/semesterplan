import openpyxl
from datetime import date

DAYS = ["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag"]


def read_schedule(path="Stundenplan.xlsx"):
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]

    try:
        col_modul  = headers.index("Modul")
        col_ects   = headers.index("ECTS")
        col_exam   = next(i for i, h in enumerate(headers) if "Pruefung" in h)
        col_diff   = next(i for i, h in enumerate(headers) if "Schwierigkeit" in h)
        day_cols   = {day: headers.index(day) for day in DAYS if day in headers}
    except (ValueError, StopIteration) as e:
        raise ValueError(f"Stundenplan.xlsx hat unerwartete Spalten: {e}")

    modules = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[col_modul]:
            continue

        sessions = []
        for day, col in day_cols.items():
            value = str(row[col]).strip() if row[col] else ""
            if value and value != "None":
                for stype in [v.strip() for v in value.split(",") if v.strip()]:
                    sessions.append({"type": stype, "day": day})

        modules.append({
            "name":       str(row[col_modul]).strip(),
            "ects":       int(row[col_ects] or 0),
            "exam_date":  str(row[col_exam]).strip() if row[col_exam] else "",
            "difficulty": int(row[col_diff] or 3),
            "start_week": date.today().isocalendar()[1],
            "sessions":   sessions,
        })

    return modules
